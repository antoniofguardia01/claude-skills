#!/usr/bin/env python3
"""Extrae todo el contenido de un .pptx para reconstruirlo como slides web.

Uso:
    python extract_pptx.py <archivo.pptx> <carpeta_salida> [--previews]

Genera en <carpeta_salida>:
    deck.json     estructura completa (textos, niveles, negritas, enlaces, enlaces
                  internos entre slides, imágenes, tablas, colores, posiciones en %)
    deck.md       resumen compacto y legible (léelo primero: ahorra tokens)
    img/          imágenes extraídas (sN-k.ext) + logos de los layouts (layout-*.ext)
    previews/     (con --previews) PNG de cada slide, para ver el diseño original
"""
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.exit("Falta python-pptx: pip install python-pptx")

WEB_OK = {"png", "jpg", "jpeg", "gif", "svg", "webp"}


def pct(v, total):
    return round((v or 0) / total * 100, 2)


def color_of(fill):
    try:
        if fill.type == 1:  # solid
            return "#" + str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def save_image(image, out_img, stem):
    ext = (image.ext or "bin").lower()
    data = image.blob
    path = out_img / f"{stem}.{ext}"
    path.write_bytes(data)
    if ext not in WEB_OK:
        # tiff/bmp/emf/wmf -> intenta png con Pillow (emf/wmf solo en Windows)
        try:
            from PIL import Image
            png = out_img / f"{stem}.png"
            Image.open(path).save(png)
            path.unlink()
            path = png
        except Exception:
            return None  # se recorta del preview más tarde
    return path.name


def paragraphs_of(tf):
    pars = []
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs) or p.text
        if not text.strip():
            continue
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            run = {"t": r.text}
            if r.font.bold:
                run["b"] = True
            if r.font.italic:
                run["i"] = True
            try:
                if r.hyperlink and r.hyperlink.address:
                    run["href"] = r.hyperlink.address
            except Exception:
                pass
            try:
                if r.font.color and r.font.color.type is not None and r.font.color.rgb:
                    run["color"] = "#" + str(r.font.color.rgb)
            except Exception:
                pass
            runs.append(run)
        size = None
        for r in p.runs:
            if r.font.size:
                size = round(r.font.size.pt)
                break
        pars.append({"text": text.strip(), "level": p.level, "size": size, "runs": runs})
    return pars


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src = Path(sys.argv[1]).resolve()
    out = Path(sys.argv[2]).resolve()
    previews = "--previews" in sys.argv
    out_img = out / "img"
    out_img.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(src))
    W, H = prs.slide_width, prs.slide_height
    slide_ids = {s.slide_id: i + 1 for i, s in enumerate(prs.slides)}
    part_to_num = {s.part.partname: i + 1 for i, s in enumerate(prs.slides)}

    deck = {"source": src.name, "ratio": round(W / H, 4), "slides": [], "layout_images": []}

    # Imágenes de layouts/master (logos, marcas de agua)
    seen_layouts = set()
    for i, s in enumerate(prs.slides, 1):
        for holder in (s.slide_layout, s.slide_layout.slide_master):
            key = holder.part.partname
            if key in seen_layouts:
                continue
            seen_layouts.add(key)
            for sh in holder.shapes:
                if getattr(sh, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        name = save_image(sh.image, out_img, f"layout-{len(deck['layout_images']) + 1}")
                        deck["layout_images"].append({"src": f"img/{name}", "from": str(key),
                            "x": pct(sh.left, W), "y": pct(sh.top, H), "w": pct(sh.width, W), "h": pct(sh.height, H)})
                    except Exception:
                        pass

    for idx, slide in enumerate(prs.slides, 1):
        items = []
        counter = [0]

        def walk(shapes, group=None):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes, group=sh.name)
                    continue
                it = {"name": sh.name, "x": pct(sh.left, W), "y": pct(sh.top, H),
                      "w": pct(sh.width, W), "h": pct(sh.height, H)}
                if group:
                    it["group"] = group
                # enlace de la forma completa
                try:
                    ca = sh.click_action
                    if ca.target_slide is not None:
                        it["goto"] = part_to_num.get(ca.target_slide.part.partname)
                    elif ca.hyperlink and ca.hyperlink.address:
                        it["href"] = ca.hyperlink.address
                except Exception:
                    pass
                image = None
                try:
                    image = sh.image  # Picture y placeholders con imagen
                except Exception:
                    image = None
                is_ole = sh.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT)
                is_frame = sh.shape_type == MSO_SHAPE_TYPE.IGX_GRAPHIC or (
                    sh.__class__.__name__ == "GraphicFrame" and not getattr(sh, "has_table", False)
                    and not getattr(sh, "has_chart", False))
                if is_ole or is_frame:
                    # objeto incrustado (Excel, SmartArt, EMF...): se recorta del render
                    counter[0] += 1
                    it["kind"] = "image"
                    it["src"] = f"img/s{idx}-{counter[0]}.png"
                    it["needs_crop"] = True
                    it["note"] = "OLE/SmartArt recortado del render de PowerPoint"
                elif image is not None:
                    counter[0] += 1
                    it["kind"] = "image"
                    name = save_image(image, out_img, f"s{idx}-{counter[0]}")
                    if name is None:
                        name = f"s{idx}-{counter[0]}.png"
                        it["needs_crop"] = True
                    it["src"] = "img/" + name
                    try:
                        it["px"] = list(image.size)
                    except Exception:
                        pass
                    # iconos diminutos (p. ej. "indicador pulsante" de enlace) = decorativos
                    if it["w"] < 3 and it["h"] < 5:
                        it["decorative"] = True
                elif getattr(sh, "has_table", False) and sh.has_table:
                    it["kind"] = "table"
                    it["rows"] = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                elif getattr(sh, "has_chart", False) and sh.has_chart:
                    it["kind"] = "chart"
                    ch = sh.chart
                    try:
                        it["chart_type"] = str(ch.chart_type)
                        it["categories"] = [str(c) for c in ch.plots[0].categories]
                        it["series"] = [{"name": s.name, "values": list(s.values)} for s in ch.series]
                    except Exception:
                        pass
                elif sh.has_text_frame and sh.text_frame.text.strip():
                    it["kind"] = "text"
                    it["paragraphs"] = paragraphs_of(sh.text_frame)
                    try:
                        if sh.is_placeholder:
                            it["placeholder"] = str(sh.placeholder_format.type).split(".")[-1].split(" ")[0]
                    except Exception:
                        pass
                else:
                    it["kind"] = "shape"
                try:
                    c = color_of(sh.fill)
                    if c:
                        it["fill"] = c
                except Exception:
                    pass
                if it["kind"] == "shape" and "fill" not in it and "goto" not in it and "href" not in it:
                    continue  # forma vacía sin valor
                items.append(it)

        walk(slide.shapes)
        # descarta número de slide y pies repetidos
        clean = []
        for it in items:
            if it.get("placeholder") in ("SLIDE_NUMBER", "FOOTER", "DATE"):
                continue
            clean.append(it)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        titles = [it for it in clean if it.get("placeholder") in ("TITLE", "CENTER_TITLE")]
        title = titles[0]["paragraphs"][0]["text"] if titles and titles[0].get("paragraphs") else None
        if not title:
            texts = [it for it in clean if it["kind"] == "text"]
            texts.sort(key=lambda t: -(t["paragraphs"][0].get("size") or 0))
            title = texts[0]["paragraphs"][0]["text"] if texts else None
        deck["slides"].append({"n": idx, "layout": slide.slide_layout.name, "title": title,
                               "items": clean, "notes": notes})

    (out / "deck.json").write_text(json.dumps(deck, ensure_ascii=False, indent=1), encoding="utf-8")

    # Resumen compacto
    md = [f"# {src.name} — {len(deck['slides'])} slides (ratio {deck['ratio']})", ""]
    if deck["layout_images"]:
        md.append("Imágenes de layout/master: " + ", ".join(li["src"] for li in deck["layout_images"]))
        md.append("")
    for s in deck["slides"]:
        md.append(f"## {s['n']}. {s['title'] or '(sin título)'}  _[{s['layout']}]_")
        for it in s["items"]:
            pos = f"@{it['x']:.0f},{it['y']:.0f} {it['w']:.0f}x{it['h']:.0f}"
            link = ""
            if it.get("goto"):
                link = f" → slide {it['goto']}"
            elif it.get("href"):
                link = f" → {it['href'][:90]}"
            if it["kind"] == "text":
                fill = f" fill={it['fill']}" if it.get("fill") else ""
                for p in it["paragraphs"]:
                    extra = ""
                    hrefs = [r["href"] for r in p["runs"] if r.get("href")]
                    if hrefs:
                        extra = f" (link: {hrefs[0][:90]})"
                    bold = "**" if p["runs"] and all(r.get("b") for r in p["runs"]) else ""
                    md.append(f"- {'  ' * p['level']}{bold}{p['text']}{bold}{extra}")
                if link or fill:
                    md.append(f"  <!-- {pos}{fill}{link} -->")
            elif it["kind"] == "image":
                if it.get("decorative"):
                    md.append(f"- (icono decorativo {it['src']}{link})")
                else:
                    md.append(f"- ![img]({it['src']}) {pos} px={it.get('px')}{link}")
            elif it["kind"] == "table":
                md.append("- TABLA:")
                for r in it["rows"]:
                    md.append("  | " + " | ".join(r) + " |")
            elif it["kind"] == "chart":
                md.append(f"- GRÁFICO {it.get('chart_type')} cats={it.get('categories')} series={it.get('series')}")
            elif it["kind"] == "shape" and (link or it.get("fill")):
                md.append(f"- forma {it.get('fill', '')} {pos}{link}")
        if s["notes"]:
            md.append(f"> Notas: {s['notes']}")
        md.append("")
    (out / "deck.md").write_text("\n".join(md), encoding="utf-8")
    print(f"OK: {len(deck['slides'])} slides -> {out / 'deck.md'}")

    pending = [(s["n"], it) for s in deck["slides"] for it in s["items"] if it.get("needs_crop")]
    if previews or pending:
        ok = render_previews(src, out / "previews")
        if pending:
            crop_pending(pending, out, ok)
            (out / "deck.json").write_text(json.dumps(deck, ensure_ascii=False, indent=1), encoding="utf-8")


def crop_pending(pending, out, ok):
    try:
        from PIL import Image
    except ImportError:
        ok = False
    for n, it in pending:
        prev = out / "previews" / f"{n:02d}.png"
        if not ok or not prev.exists():
            print(f"  ! slide {n}: {it['src']} necesita recorte manual (sin render)")
            continue
        im = Image.open(prev)
        W, H = im.size
        box = (int(it["x"] / 100 * W), int(it["y"] / 100 * H),
               int((it["x"] + it["w"]) / 100 * W), int((it["y"] + it["h"]) / 100 * H))
        im.crop(box).save(out / it["src"])
        del it["needs_crop"]
        print(f"  recortado {it['src']} del render de la slide {n}")


def render_previews(src, dest):
    dest.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        ps = (f"$pp=New-Object -ComObject PowerPoint.Application;"
              f"$p=$pp.Presentations.Open('{src}',$true,$false,$false);"
              f"$p.Export('{dest}','PNG',2560,1440);$p.Close();$pp.Quit()")
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
        if r.returncode == 0 and any(dest.iterdir()):
            for f in dest.iterdir():  # Diapositiva12.PNG / Slide12.PNG -> 12.png
                num = "".join(ch for ch in f.stem if ch.isdigit())
                if num:
                    f.rename(dest / f"{int(num):02d}.png")
            print(f"Previews (PowerPoint, 2560px) -> {dest}")
            return True
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(dest), str(src)], capture_output=True)
        pdf = dest / (src.stem + ".pdf")
        try:
            import pymupdf
            doc = pymupdf.open(pdf)
            for i, page in enumerate(doc, 1):
                page.get_pixmap(dpi=200).save(dest / f"{i:02d}.png")
            print(f"Previews (LibreOffice) -> {dest}")
            return True
        except Exception:
            pass
    print("! No se pudieron generar previews (ni PowerPoint ni LibreOffice). Trabaja con deck.md.")
    return False


if __name__ == "__main__":
    main()
